import csv
import pdfkit
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill
from django.conf import settings
from django.http import HttpResponse
from django.template.loader import render_to_string
from django.shortcuts import get_object_or_404
from django.contrib.auth.decorators import login_required
from django.utils import timezone

from ..models import BillOfMaterials

@login_required
def export_bom_pdf(request, bom_id):
    """
    Exports a single Bill of Materials to a PDF file.
    """
    bom = get_object_or_404(BillOfMaterials.objects.select_related(
        'product', 'size_group', 'created_by'
    ).prefetch_related('items__material__unit_new'), pk=bom_id)
    
    # Render HTML template to a string
    html = render_to_string('pdf/production/bom_pdf.html', {'bom': bom})
    
    # Configure PDF options
    options = {
        'page-size': 'A4',
        'margin-top': '0.75in',
        'margin-right': '0.75in',
        'margin-bottom': '0.75in',
        'margin-left': '0.75in',
        'encoding': "UTF-8",
        '--header-font-name': 'Tajawal',
        '--footer-font-name': 'Tajawal',
        '--load-error-handling': 'ignore',
        '--load-media-error-handling': 'ignore',
    }

    try:
        # Generate PDF from HTML string
        pdf = pdfkit.from_string(html, False, options=options, configuration=pdfkit.configuration(wkhtmltopdf=settings.WKHTMLTOPDF_PATH))
        
        # Create HTTP response
        response = HttpResponse(pdf, content_type='application/pdf')
        filename = f"BOM_{bom.product.code}_{bom.version}_{timezone.now().strftime('%Y%m%d')}.pdf"
        response['Content-Disposition'] = f'attachment; filename="{filename}"'
        
        return response
    except Exception as e:
        # Handle exceptions, e.g., wkhtmltopdf not found
        return HttpResponse(f"Error generating PDF: {e}", status=500)

@login_required
def export_bom_csv(request, bom_id):
    """
    Exports the items of a Bill of Materials to a CSV file.
    """
    bom = get_object_or_404(BillOfMaterials, pk=bom_id)
    
    response = HttpResponse(content_type='text/csv; charset=utf-8-sig') # utf-8-sig adds BOM for Excel
    filename = f"BOM_{bom.product.code}_{bom.version}_{timezone.now().strftime('%Y%m%d')}.csv"
    response['Content-Disposition'] = f'attachment; filename="{filename}"'

    writer = csv.writer(response)
    
    # Write headers
    writer.writerow([
        'المادة الخام', 'كود المادة', 'الكمية', 'الوحدة', 'تكلفة الوحدة', 'التكلفة الإجمالية', 'ملاحظات'
    ])

    # Write data rows
    for item in bom.items.all():
        writer.writerow([
            item.material.name,
            item.material.code,
            item.quantity,
            item.material.unit_new.symbol if item.material.unit_new else 'وحدة',
            item.material.cost_price,
            item.cost,
            item.notes
        ])
    
    # Write total cost row
    writer.writerow([])
    writer.writerow(['', '', '', '', 'الإجمالي', bom.total_cost, ''])

    return response

@login_required
def export_bom_excel(request, bom_id):
    """
    Exports the items of a Bill of Materials to an Excel file.
    """
    bom = get_object_or_404(BillOfMaterials, pk=bom_id)
    
    # Create workbook and sheet
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = f"BOM v{bom.version}"
    ws.sheet_view.rightToLeft = True

    # Define styles
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="4F81BD", end_color="4F81BD", fill_type="solid")
    header_alignment = Alignment(horizontal="center", vertical="center")
    
    # Write headers
    headers = ['المادة الخام', 'كود المادة', 'الكمية', 'الوحدة', 'تكلفة الوحدة', 'التكلفة الإجمالية', 'ملاحظات']
    for col_num, header_title in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_num, value=header_title)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_alignment

    # Write data
    for row_num, item in enumerate(bom.items.all(), 2):
        ws.cell(row=row_num, column=1, value=item.material.name)
        ws.cell(row=row_num, column=2, value=item.material.code)
        ws.cell(row=row_num, column=3, value=item.quantity).number_format = '0.000'
        ws.cell(row=row_num, column=4, value=item.material.unit_new.symbol if item.material.unit_new else 'وحدة')
        ws.cell(row=row_num, column=5, value=item.material.cost_price).number_format = '#,##0.00'
        ws.cell(row=row_num, column=6, value=item.cost).number_format = '#,##0.00'
        ws.cell(row=row_num, column=7, value=item.notes)
    
    # Add Total row
    total_row = ws.max_row + 2
    total_cell = ws.cell(row=total_row, column=5, value="الإجمالي")
    total_cell.font = Font(bold=True)
    ws.cell(row=total_row, column=6, value=bom.total_cost).number_format = '#,##0.00'
    ws.cell(row=total_row, column=6).font = Font(bold=True)
    
    # Adjust column widths
    ws.column_dimensions['A'].width = 35
    ws.column_dimensions['B'].width = 15
    ws.column_dimensions['C'].width = 12
    ws.column_dimensions['D'].width = 10
    ws.column_dimensions['E'].width = 15
    ws.column_dimensions['F'].width = 18
    ws.column_dimensions['G'].width = 40

    # Create response
    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    filename = f"BOM_{bom.product.code}_{bom.version}_{timezone.now().strftime('%Y%m%d')}.xlsx"
    response['Content-Disposition'] = f'attachment; filename="{filename}"'
    
    wb.save(response)
    return response

