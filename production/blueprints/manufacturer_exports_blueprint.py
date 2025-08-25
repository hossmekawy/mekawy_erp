import openpyxl
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
from openpyxl.utils import get_column_letter
from django.http import HttpResponse
from django.shortcuts import get_object_or_404, reverse
from django.template.loader import render_to_string
from django.conf import settings
from django.urls import path
from django.contrib.auth.decorators import login_required
from django.utils import timezone
import pdfkit
from decimal import Decimal
import logging
from itertools import chain
from collections import defaultdict

from ..models import (
    ExternalManufacturer, AssemblyProcess, DyeingProcess, FinishingProcess,
    AssemblyComponent, FinishingComponent
)
from ..templatetags.production_extras import humanize_duration

logger = logging.getLogger(__name__)

def get_full_manufacturer_context(pk):
    """
    Helper function to gather all data related to a manufacturer for reporting.
    """
    manufacturer = get_object_or_404(
        ExternalManufacturer.objects.prefetch_related(
            'assembly_processes__production_order__product',
            'dyeing_jobs__assembly_process__production_order__product',
            'finishing_jobs__dyeing_process__assembly_process__production_order__product',
            'product_prices__product'
        ).select_related('finance_account'), 
        pk=pk
    )

    # --- Combine all job types into a single list ---
    assembly_jobs = manufacturer.assembly_processes.all()
    dyeing_jobs = manufacturer.dyeing_jobs.all()
    finishing_jobs = manufacturer.finishing_jobs.all()
    all_jobs = list(chain(assembly_jobs, dyeing_jobs, finishing_jobs))
    active_jobs_raw = [job for job in all_jobs if not job.is_completed]
    completed_jobs_raw = [job for job in all_jobs if job.is_completed]
    default_date = timezone.now()

    # --- Standardize the Active Jobs list ---
    active_jobs_list = []
    for job in active_jobs_raw:
        production_order, job_type_display, start_date, detail_url, quantity = None, "غير محدد", None, "#", 0
        if isinstance(job, AssemblyProcess):
            production_order, job_type_display, start_date, detail_url, quantity = job.production_order, "تجميع", job.start_date, reverse('production:assembly_detail', kwargs={'pk': job.pk}), job.quantity_sent
        elif isinstance(job, DyeingProcess):
            production_order, job_type_display, start_date, detail_url, quantity = job.assembly_process.production_order, "صباغة", job.sent_date, reverse('production:dyeing_detail', kwargs={'pk': job.pk}), job.quantity_sent
        elif isinstance(job, FinishingProcess):
            production_order, job_type_display, start_date, detail_url, quantity = job.dyeing_process.assembly_process.production_order, "تشطيب", job.start_date, reverse('production:finishing_detail', kwargs={'pk': job.pk}), job.quantity_input
        if production_order:
            active_jobs_list.append({'order': production_order, 'job_type': job_type_display, 'start_date': start_date, 'detail_url': detail_url, 'quantity': quantity})
    
    # --- Standardize the Job History list ---
    job_history_list = []
    total_completed_value = Decimal('0.0')
    for job in completed_jobs_raw:
        cost, completion_date, detail_url, production_order, job_type_display, quantity = 0, None, "#", None, "غير محدد", 0
        if isinstance(job, AssemblyProcess):
            cost, completion_date, production_order, job_type_display, detail_url, quantity = job.assembly_cost, job.actual_completion_date, job.production_order, "تجميع", reverse('production:assembly_detail', kwargs={'pk': job.pk}), job.quantity_received
        elif isinstance(job, DyeingProcess):
            cost, completion_date, production_order, job_type_display, detail_url, quantity = job.total_dyeing_cost, job.actual_return_date, job.assembly_process.production_order, "صباغة", reverse('production:dyeing_detail', kwargs={'pk': job.pk}), job.quantity_received
        elif isinstance(job, FinishingProcess):
            cost, completion_date, production_order, job_type_display, detail_url, quantity = job.total_finishing_cost, job.actual_completion_date, job.dyeing_process.assembly_process.production_order, "تشطيب", reverse('production:finishing_detail', kwargs={'pk': job.pk}), job.quantity_output
        if production_order:
            job_history_list.append({'order': production_order, 'cost': cost or 0, 'completion_date': completion_date, 'job_type': job_type_display, 'detail_url': detail_url, 'quantity': quantity})
            total_completed_value += cost or 0

    # --- Aggregate All Sent Materials ---
    materials_by_order = defaultdict(list)
    assembly_components = AssemblyComponent.objects.filter(assembly_process__external_manufacturer=manufacturer).select_related('assembly_process__production_order', 'material__unit_new', 'source_warehouse')
    for comp in assembly_components:
        materials_by_order[comp.assembly_process.production_order].append({'material': comp.material, 'quantity': comp.quantity_sent, 'stage': 'تجميع / إضافي', 'date': comp.assembly_process.start_date})
    finishing_components = FinishingComponent.objects.filter(finishing_process__external_manufacturer=manufacturer).select_related('finishing_process__dyeing_process__assembly_process__production_order', 'material__unit_new', 'source_warehouse')
    for comp in finishing_components:
        materials_by_order[comp.finishing_process.dyeing_process.assembly_process.production_order].append({'material': comp.material, 'quantity': comp.quantity_sent, 'stage': 'تشطيب', 'date': comp.finishing_process.start_date})

    return {
        'manufacturer': manufacturer,
        'active_jobs': sorted(active_jobs_list, key=lambda x: x['start_date'] or default_date, reverse=True),
        'job_history': sorted(job_history_list, key=lambda x: x['completion_date'] or default_date, reverse=True),
        'materials_by_order': dict(sorted(materials_by_order.items(), key=lambda item: max(m['date'] for m in item[1]), reverse=True)),
        'product_prices': manufacturer.product_prices.all(),
        'total_value_of_completed_jobs': total_completed_value
    }

@login_required
def export_manufacturer_pdf(request, pk):
    context = get_full_manufacturer_context(pk)
    html_string = render_to_string('pdf/production/manufacturer_report.html', context)
    try:
        pdf_config = pdfkit.configuration(wkhtmltopdf=settings.WKHTMLTOPDF_PATH)
        pdf = pdfkit.from_string(html_string, False, configuration=pdf_config, options={
            'encoding': "UTF-8", 'page-size': 'A4', 'margin-top': '0.75in',
            'margin-right': '0.75in', 'margin-bottom': '0.75in', 'margin-left': '0.75in',
        })
        response = HttpResponse(pdf, content_type='application/pdf')
        response['Content-Disposition'] = f'inline; filename="Manufacturer_{context["manufacturer"].name}.pdf"'
        return response
    except Exception as e:
        logger.error(f"PDF Generation Error for Manufacturer {pk}: {e}")
        return HttpResponse(f"Error generating PDF: {e}", status=500)

@login_required
def export_manufacturer_excel(request, pk):
    context = get_full_manufacturer_context(pk)
    manufacturer = context['manufacturer']
    wb = openpyxl.Workbook()

    # --- Define Styles ---
    title_font = Font(name='Calibri', size=18, bold=True, color="FFFFFF")
    header_font = Font(name='Calibri', size=12, bold=True, color="FFFFFF")
    section_font = Font(name='Calibri', size=13, bold=True)
    cell_font = Font(name='Calibri', size=11)
    
    header_fill = PatternFill(start_color="4A86E8", end_color="4A86E8", fill_type="solid")
    center_align = Alignment(horizontal='center', vertical='center', wrap_text=True)

    # --- Summary Sheet ---
    ws = wb.active
    ws.title = "ملخص المصنع"
    ws.sheet_view.rightToLeft = True
    ws.merge_cells('A1:E2')
    ws['A1'] = f'تقرير المصنع: {manufacturer.name}'
    ws['A1'].font = title_font
    ws['A1'].fill = header_fill
    ws['A1'].alignment = center_align

    info = [
        ('الشخص المسؤول', manufacturer.contact_person),
        ('الهاتف', manufacturer.phone),
        ('العنوان', manufacturer.address),
        ('الرصيد المالي', f"{manufacturer.finance_account.balance if manufacturer.finance_account else 0:,.2f} جنيه"),
        ('قيمة الأعمال النشطة', f"{manufacturer.total_value_of_active_jobs:,.2f} جنيه"),
        ('قيمة الأعمال المكتملة', f"{context['total_value_of_completed_jobs']:,.2f} جنيه"),
    ]
    for i, (label, value) in enumerate(info, 4):
        ws[f'A{i}'] = label
        ws[f'B{i}'] = value
        ws[f'A{i}'].font = Font(bold=True)
    
    # --- Active Jobs Sheet ---
    ws2 = wb.create_sheet(title="المهام النشطة")
    ws2.sheet_view.rightToLeft = True
    headers = ['أمر الإنتاج', 'المنتج', 'الكمية', 'نوع المهمة', 'تاريخ الإرسال']
    ws2.append(headers)
    for cell in ws2[1]:
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center_align
    for job in context['active_jobs']:
        ws2.append([
            job['order'].order_number, job['order'].product.name, job['quantity'],
            job['job_type'], job['start_date'].strftime('%Y-%m-%d') if job['start_date'] else ''
        ])

    # --- Job History Sheet ---
    ws3 = wb.create_sheet(title="سجل المهام")
    ws3.sheet_view.rightToLeft = True
    headers = ['أمر الإنتاج', 'المنتج', 'الكمية', 'نوع المهمة', 'التكلفة (جنيه)', 'تاريخ الإكمال']
    ws3.append(headers)
    for cell in ws3[1]:
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center_align
    for job in context['job_history']:
        ws3.append([
            job['order'].order_number, job['order'].product.name, job['quantity'],
            job['job_type'], job['cost'], job['completion_date'].strftime('%Y-%m-%d') if job['completion_date'] else ''
        ])
        ws3.cell(row=ws3.max_row, column=5).number_format = '#,##0.00'

    # --- Materials Sheet ---
    ws4 = wb.create_sheet(title="سجل المواد المرسلة")
    ws4.sheet_view.rightToLeft = True
    headers = ['أمر الإنتاج', 'المادة الخام', 'الكمية', 'الوحدة', 'المرحلة', 'تاريخ الإرسال']
    ws4.append(headers)
    for cell in ws4[1]:
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center_align
    for order, materials in context['materials_by_order'].items():
        for mat_data in materials:
            ws4.append([
                order.order_number, mat_data['material'].name, mat_data['quantity'],
                mat_data['material'].unit_new.symbol if mat_data['material'].unit_new else 'قطعة',
                mat_data['stage'], mat_data['date'].strftime('%Y-%m-%d') if mat_data['date'] else ''
            ])

    # Auto-size columns for all sheets
    for sheet in wb.worksheets:
        for col in sheet.columns:
            sheet.column_dimensions[get_column_letter(col[0].column)].auto_size = True

    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = f'attachment; filename="Manufacturer_Report_{manufacturer.name}.xlsx"'
    wb.save(response)
    return response

urlpatterns = [
    path('pdf/', export_manufacturer_pdf, name='pdf'),
    path('excel/', export_manufacturer_excel, name='excel'),
]
