import openpyxl
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
from openpyxl.utils import get_column_letter
from django.http import HttpResponse
from django.shortcuts import get_object_or_404
from django.template.loader import render_to_string
from django.conf import settings
from django.urls import path
from django.contrib.auth.decorators import login_required
import pdfkit
from decimal import Decimal
import logging

from ..models import (
    ProductionCostAnalysis, AssemblyComponent, FinishingComponent, 
    BOMItem, DyeingProcess, ManufacturerProductPrice, ProductionOrder,
    AssemblyProcess, FinishingProcess
)
from warehouses.models import Product, ProductBatch
from ..templatetags.production_extras import humanize_duration

logger = logging.getLogger(__name__)

def get_full_context(analysis_id):
    """
    Helper function to gather all data for the reports.
    This function is now updated with the corrected database query.
    """
    analysis = get_object_or_404(
        ProductionCostAnalysis.objects.select_related(
            'production_order__product',
            # --- FIX: Correct path to the size_group through the bom_version ---
            'production_order__bom_version__size_group', 
            'production_order__textile_stock__product', 
            'production_order__cutting_process',
        ), 
        pk=analysis_id
    )
    order = analysis.production_order
    bom = order.bom_version
    cutting_process = getattr(order, 'cutting_process', None)

    # --- Process Timeline & Durations (Logic updated to match detail view) ---
    timeline = []
    if cutting_process and cutting_process.cutting_date and cutting_process.completed_at:
        duration = cutting_process.completed_at - cutting_process.cutting_date
        timeline.append({
            'stage': 'القص', 'handler': cutting_process.cutter.get_full_name() if cutting_process.cutter else 'داخلي',
            'start': cutting_process.cutting_date, 'end': cutting_process.completed_at, 'duration': duration,
            'cost_per_piece': None
        })

    for assembly in order.assembly_processes.all():
        duration = (assembly.actual_completion_date - assembly.start_date) if assembly.start_date and assembly.actual_completion_date else None
        handler = assembly.external_manufacturer.name if assembly.assembly_type == 'outsourced' and assembly.external_manufacturer else (assembly.assembler.get_full_name() if assembly.assembler else 'داخلي')
        cost_per_piece = None
        if assembly.assembly_type == 'outsourced' and assembly.external_manufacturer:
            price_obj = ManufacturerProductPrice.objects.filter(manufacturer=assembly.external_manufacturer, product=order.product).first()
            cost_per_piece = price_obj.price if price_obj else None
        timeline.append({
            'stage': f'التجميع ({assembly.get_assembly_type_display()})', 'start': assembly.start_date, 'end': assembly.actual_completion_date, 
            'duration': duration, 'handler': handler, 'cost_per_piece': cost_per_piece
        })

        for dyeing in assembly.dyeing_processes.all():
            duration = (dyeing.actual_return_date - dyeing.sent_date) if dyeing.sent_date and dyeing.actual_return_date else None
            handler = dyeing.dyeing_facility.name if dyeing.dyeing_facility else 'غير محدد'
            cost_per_piece_dyeing = None
            if dyeing.dyeing_facility:
                price_obj = ManufacturerProductPrice.objects.filter(manufacturer=dyeing.dyeing_facility, product=order.product).first()
                cost_per_piece_dyeing = price_obj.price if price_obj else None
            timeline.append({
                'stage': f'الصباغة ({dyeing.color_specification})', 'start': dyeing.sent_date, 'end': dyeing.actual_return_date, 
                'duration': duration, 'handler': handler, 'cost_per_piece': cost_per_piece_dyeing
            })
            if finishing_process := getattr(dyeing, 'finishing_process', None):
                duration = (finishing_process.actual_completion_date - finishing_process.start_date) if finishing_process.start_date and finishing_process.actual_completion_date else None
                handler = finishing_process.external_manufacturer.name if finishing_process.finishing_type == 'outsourced' and finishing_process.external_manufacturer else (finishing_process.finisher.get_full_name() if finishing_process.finisher else 'داخلي')
                cost_per_piece_finishing = None
                if finishing_process.finishing_type == 'outsourced' and finishing_process.external_manufacturer:
                    price_obj = ManufacturerProductPrice.objects.filter(manufacturer=finishing_process.external_manufacturer, product=order.product).first()
                    cost_per_piece_finishing = price_obj.price if price_obj else None
                timeline.append({
                    'stage': 'التشطيب', 'start': finishing_process.start_date, 'end': finishing_process.actual_completion_date, 
                    'duration': duration, 'handler': handler, 'cost_per_piece': cost_per_piece_finishing
                })

    # --- BOM vs Actual Usage ---
    bom_comparison = []
    if bom:
        bom_materials = {item.material_id: item for item in bom.items.all()}
        actual_assembly = {c.material_id: c.quantity_sent for c in AssemblyComponent.objects.filter(assembly_process__production_order=order)}
        actual_finishing = {c.material_id: c.quantity_sent for c in FinishingComponent.objects.filter(finishing_process__dyeing_process__assembly_process__production_order=order)}
        actual_materials = actual_assembly
        for mat_id, qty in actual_finishing.items():
            actual_materials[mat_id] = actual_materials.get(mat_id, 0) + qty
        all_material_ids = set(bom_materials.keys()) | set(actual_materials.keys())
        for mat_id in all_material_ids:
            bom_item = bom_materials.get(mat_id)
            actual_qty = actual_materials.get(mat_id, Decimal('0.0'))
            material_obj = bom_item.material if bom_item else Product.objects.get(id=mat_id)
            bom_qty_per_piece = bom_item.quantity if bom_item else Decimal('0.0')
            bom_total_qty = bom_qty_per_piece * order.quantity_ordered
            bom_comparison.append({
                'material': material_obj.name,
                'bom_quantity': bom_total_qty,
                'actual_quantity': actual_qty,
                'variance': actual_qty - bom_total_qty
            })
            
    # --- Batch Data dictionary ---
    batch_data = {
        'order': order,
        'cutting_process': cutting_process,
        'colors': list(set(DyeingProcess.objects.filter(assembly_process__production_order=order).values_list('color_specification', flat=True))),
        'storage_locations': list(set(ProductBatch.objects.filter(
            stock__product=order.product,
            batch_number=order.batch_number
        ).values_list('stock__warehouse__name', flat=True)))
    }
    
    fabric_bom_quantity = None
    if bom and order.textile_stock:
        try:
            fabric_material = order.textile_stock.product
            fabric_bom_item = BOMItem.objects.get(bom=bom, material=fabric_material)
            fabric_bom_quantity = fabric_bom_item.quantity
        except BOMItem.DoesNotExist:
            fabric_bom_quantity = None
            
    batch_data['fabric_bom_quantity'] = fabric_bom_quantity
    batch_data['fabric_cutting_quantity'] = cutting_process.single_layer_meterage if cutting_process else None

    return {
        'analysis': analysis,
        'order': order,
        'timeline': timeline,
        'bom_comparison': bom_comparison,
        'batch_data': batch_data,
    }

@login_required
def export_cost_analysis_pdf(request, pk):
    context = get_full_context(pk)
    html_string = render_to_string('pdf/production/cost_analysis_report.html', context)
    try:
        pdf_config = pdfkit.configuration(wkhtmltopdf=settings.WKHTMLTOPDF_PATH)
        pdf = pdfkit.from_string(html_string, False, configuration=pdf_config, options={
            'encoding': "UTF-8", 'page-size': 'A4', 'margin-top': '0.5in',
            'margin-right': '0.5in', 'margin-bottom': '0.5in', 'margin-left': '0.5in',
        })
        response = HttpResponse(pdf, content_type='application/pdf')
        response['Content-Disposition'] = f'attachment; filename="CostAnalysis_{context["order"].order_number}.pdf"'
        return response
    except Exception as e:
        logging.error(f"PDF Generation Error: {e}")
        return HttpResponse(f"Error generating PDF: {e}", status=500)

@login_required
def export_cost_analysis_excel(request, pk):
    """
    Exports a fully-formatted, professional Excel report for the Cost Analysis.
    """
    context = get_full_context(pk)
    analysis = context['analysis']
    order = context['order']
    
    wb = openpyxl.Workbook()
    
    # --- Define Styles ---
    title_font = Font(name='Calibri', size=18, bold=True, color="FFFFFF")
    header_font = Font(name='Calibri', size=12, bold=True, color="FFFFFF")
    section_font = Font(name='Calibri', size=13, bold=True)
    cell_font = Font(name='Calibri', size=11)
    total_font = Font(name='Calibri', size=11, bold=True)
    
    header_fill = PatternFill(start_color="4A86E8", end_color="4A86E8", fill_type="solid") # A modern blue
    section_fill = PatternFill(start_color="D9EAD3", end_color="D9EAD3", fill_type="solid") # A light green
    total_fill = PatternFill(start_color="F3F3F3", end_color="F3F3F3", fill_type="solid")

    center_align = Alignment(horizontal='center', vertical='center')
    right_align = Alignment(horizontal='right', vertical='center')
    
    thin_border = Border(left=Side(style='thin'), right=Side(style='thin'), top=Side(style='thin'), bottom=Side(style='thin'))

    # --- Summary Sheet ---
    ws = wb.active
    ws.title = "ملخص التقرير"
    ws.sheet_view.rightToLeft = True

    # Main Title
    ws.merge_cells('A1:E2')
    ws['A1'] = f'تقرير تحليل التكلفة: {order.order_number}'
    ws['A1'].font = title_font
    ws['A1'].fill = header_fill
    ws['A1'].alignment = center_align

    # Order Info
    ws.append([]) # Spacer
    ws.append(['المنتج:', order.product.name, '', 'الكمية:', f'{order.quantity_ordered} قطعة'])
    ws['A3'].font = total_font
    ws['D3'].font = total_font
    
    # Metrics
    ws.append([]) # Spacer
    ws.append(['إجمالي تكلفة الإنتاج', 'تكلفة القطعة', 'نسبة المواد', 'نسبة العمليات'])
    for cell in ws[5]:
        cell.font = header_font
        cell.fill = section_fill
        cell.alignment = center_align
    
    ws.append([
        analysis.total_production_cost,
        analysis.cost_per_piece,
        analysis.material_cost_percentage / 100,
        analysis.process_cost_percentage / 100
    ])
    ws['A6'].number_format = '#,##0.00 "جنيه"'
    ws['B6'].number_format = '#,##0.00 "جنيه"'
    ws['C6'].number_format = '0.00%'
    ws['D6'].number_format = '0.00%'
    for cell in ws[6]:
        cell.font = cell_font
        cell.alignment = center_align

    # Cost Breakdown Table
    ws.append([]) # Spacer
    ws.append(['تفاصيل التكاليف', 'التكلفة (جنيه)'])
    ws['A8'].font = header_font
    ws['B8'].font = header_font
    ws['A8'].fill = section_fill
    ws['B8'].fill = section_fill
    
    cost_data = [
        ('تكاليف المواد الخام', None, True),
        ('تكلفة القماش', analysis.fabric_cost),
        ('تكلفة الخيوط', analysis.thread_cost),
        ('تكلفة الإكسسوارات', analysis.accessories_cost),
        ('إجمالي تكلفة المواد', analysis.total_material_cost, False, True),
        ('تكاليف العمليات', None, True),
        ('تكلفة القص', analysis.cutting_cost),
        ('تكلفة التجميع', analysis.assembly_cost),
        ('تكلفة الصباغة', analysis.dyeing_cost),
        ('تكلفة التشطيب', analysis.finishing_cost),
        ('إجمالي تكلفة العمليات', analysis.total_process_cost, False, True),
        ('تكاليف إضافية', None, True),
        ('تكلفة النقل', analysis.transportation_cost),
        ('التكاليف العامة', analysis.overhead_cost),
        ('مراقبة الجودة', analysis.quality_control_cost),
        ('الفاقد والعيوب', analysis.waste_cost + analysis.defect_cost),
        ('إجمالي التكاليف الإضافية', analysis.total_additional_cost, False, True),
    ]

    for row_data in cost_data:
        is_section = row_data[2] if len(row_data) > 2 else False
        is_total = row_data[3] if len(row_data) > 3 else False
        
        ws.append([row_data[0], row_data[1]])
        last_row = ws.max_row
        cell1 = ws[f'A{last_row}']
        cell2 = ws[f'B{last_row}']
        
        cell1.font = cell_font
        cell2.font = cell_font
        cell2.number_format = '#,##0.00'
        
        if is_section:
            ws.merge_cells(f'A{last_row}:B{last_row}')
            cell1.fill = section_fill
            cell1.font = section_font
        if is_total:
            cell1.font = total_font
            cell2.font = total_font
            cell1.fill = total_fill
            cell2.fill = total_fill
    
    for col_letter in ['A', 'B', 'C', 'D', 'E']:
        ws.column_dimensions[col_letter].width = 25

    # --- Timeline Sheet ---
    ws2 = wb.create_sheet(title="الجدول الزمني")
    ws2.sheet_view.rightToLeft = True
    ws2.append(['المرحلة', 'الجهة المنفذة', 'تكلفة القطعة', 'تاريخ البدء', 'تاريخ الانتهاء', 'المدة'])
    for cell in ws2[1]:
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center_align

    for item in context['timeline']:
        duration_str = humanize_duration(item['duration']) if item['duration'] else ''
        ws2.append([
            item['stage'], 
            item.get('handler', ''),
            item.get('cost_per_piece', ''),
            item['start'].strftime('%Y-%m-%d %H:%M') if item['start'] else '', 
            item['end'].strftime('%Y-%m-%d %H:%M') if item['end'] else '', 
            duration_str
        ])
        if item.get('cost_per_piece'):
            ws2.cell(row=ws2.max_row, column=3).number_format = '#,##0.00'

    for col in ws2.columns:
        ws2.column_dimensions[get_column_letter(col[0].column)].auto_size = True

    # --- BOM Comparison Sheet ---
    ws3 = wb.create_sheet(title="مقارنة الخامات")
    ws3.sheet_view.rightToLeft = True
    ws3.append(['المادة الخام', 'الكمية حسب BOM', 'الكمية الفعلية', 'الفرق'])
    for cell in ws3[1]:
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center_align

    for item in context['bom_comparison']:
        ws3.append([item['material'], item['bom_quantity'], item['actual_quantity'], item['variance']])
        last_row = ws3.max_row
        for i in range(2, 5):
            ws3.cell(row=last_row, column=i).number_format = '#,##0.00'
        
    for col in ws3.columns:
        ws3.column_dimensions[get_column_letter(col[0].column)].auto_size = True
        
    for row in ws3.iter_rows(min_row=2, max_col=4, max_row=ws3.max_row):
        variance_cell = row[3]
        if variance_cell.value > 0:
            variance_cell.font = Font(color="9C0006") # Red
            variance_cell.fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
        elif variance_cell.value < 0:
            variance_cell.font = Font(color="006100") # Green
            variance_cell.fill = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")

    # Finalize and send response
    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = f'attachment; filename="CostAnalysis_{order.order_number}.xlsx"'
    wb.save(response)
    return response

urlpatterns = [
    path('pdf/', export_cost_analysis_pdf, name='pdf'),
    path('excel/', export_cost_analysis_excel, name='excel'),
]
